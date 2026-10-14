"""Permission-respecting operations dashboard read model and cycle close action."""
import json
from io import BytesIO
from html import unescape
from html.parser import HTMLParser

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import frappe

PAGE_SIZE = 500
MAX_ROWS = 20000

# Owner-confirmed Operations routing. Employee.user_id resolves the opener to
# this allowlisted roster; these names do not grant access to a Site Cycle.
OPS_ROSTER = {
    "Capex": {
        "pm": ("Mohammed Sohail Mohammed Ilyas",),
        "co": ("Ghlzar Ahmad Fazal Malik",),
    },
    "Incident": {
        "pm": ("Abdulraheem Mohammed", "Mohammed Sohail Mohammed Ilyas"),
        "co": ("Arshad ahmed abdulhameed",),
    },
    "Colocation": {
        "pm": ("Mohammed Imran Ansary",),
        "co": ("Mohamed Walid Khalil Nmer Rahma", "Obed Sharrif"),
    },
}

# Warehouse owns the operational DN closure check. These accounts are the
# confirmed operational contacts; this list does not grant Page or DocType
# access. Ramzan currently lacks SIG Site Cycle/Workbench access, so the
# permission checks below intentionally leave his queue unavailable until the
# site role setup is changed through its own reviewed release.
WAREHOUSE_USERS = {
    "m.ramzan@sigtele.com",
    "tahir.shahzad@sigtele.com",
}

# Short owner labels confirmed for the Operations dashboard. Stage work remains
# shared by its project-family role; these names explain who that role covers
# and do not assign an individual or change access.
OPS_OWNER_NAMES = {
    "Capex": {"pm": ("Sohil",), "co": ("Gulzar",)},
    "Incident": {"pm": ("Sohil", "Raheem"), "co": ("Arshad",)},
    "Colocation": {"pm": ("Imran",), "co": ("Waleed", "Obed")},
}
WAREHOUSE_OWNER_NAMES = ("Ramzan", "Tahir")

# The Site Cycles page is a cross-project dashboard for the same audiences
# allowed by its Page role list. Data reads still use normal Frappe permissions.
SITE_CYCLE_DASHBOARD_ROLES = {
    "System Manager", "Projects Manager", "Projects User", "SIG Field Ops",
}

TASK_STAGE = {
    "mapping": "1 - Initial", "boq": "1 - Initial",
    "tcn": "2 - Plan", "mir": "2 - Plan", "visit": "2 - Plan",
    "tags": "3 - Handover", "photos": "3 - Handover", "asbuilt": "3 - Handover",
    "asbuilt_operational": "3 - Handover",
}


def _normalize_person(value):
    return " ".join(str(value or "").split()).casefold()


def _current_ops_membership(user):
    """Resolve only this opener's Employee link and confirmed Operations scope."""
    if user in WAREHOUSE_USERS:
        return {"ready": True, "issue": "", "memberships": [
            {"scope": scope, "role": "warehouse"} for scope in OPS_ROSTER
        ]}
    if user in ("", "Guest") or not frappe.db.exists("DocType", "Employee"):
        return {"ready": False, "issue": "Employee identity is unavailable.", "memberships": []}
    if not frappe.has_permission("Employee", "read"):
        return {"ready": False, "issue": "Your Employee link cannot be read for task routing.", "memberships": []}
    try:
        employees = frappe.get_list(
            "Employee", filters={"user_id": user, "status": "Active"},
            fields=["name", "employee_name", "user_id"], page_length=2)
    except Exception:
        return {"ready": False, "issue": "Employee identity lookup needs review.", "memberships": []}
    if len(employees) != 1:
        return {"ready": False, "issue": "A unique active Employee link is required for task routing.", "memberships": []}
    employee_name = _normalize_person(employees[0].get("employee_name"))
    memberships = []
    for scope, roles in OPS_ROSTER.items():
        for role, names in roles.items():
            if any(_normalize_person(name) == employee_name for name in names):
                memberships.append({"scope": scope, "role": role})
    if not memberships:
        return {"ready": False, "issue": "This Employee is not in the confirmed Operations roster.", "memberships": []}
    return {"ready": True, "issue": "", "memberships": memberships}


def _task_audience(family, owner_role, memberships):
    """A shared task is visible to a roster member of its scope and role."""
    if owner_role == "either":
        return any(m["scope"] == family and m["role"] in ("pm", "co") for m in memberships)
    return any(m["scope"] == family and m["role"] == owner_role for m in memberships)


def _owner_names(family, owner_role):
    """Return short display names for the confirmed shared task audience."""
    if owner_role == "warehouse":
        return list(WAREHOUSE_OWNER_NAMES)
    family_owners = OPS_OWNER_NAMES.get(family, {})
    if owner_role == "either":
        return list(dict.fromkeys(family_owners.get("pm", ()) + family_owners.get("co", ())))
    return list(family_owners.get(owner_role, ()))


def _asbuilt_operational_status(cycle, lines):
    """Use an explicit site DN line decision; quantities never imply closure."""
    site_refs = {
        str(cycle.get(field) or "").strip().casefold()
        for field in ("site_key", "site") if str(cycle.get(field) or "").strip()
    }
    matched = []
    for line in lines:
        site = str(line.get("custom_site") or "").strip().casefold()
        if not site or site not in site_refs:
            continue
        matched.append(line)
    return {
        "applicable": bool(matched),
        "closed": any(_flag(line, "custom_return_closed") for line in matched),
        "source_dns": sorted({str(line.get("parent") or "") for line in matched if line.get("parent")}),
    }


def _can_read_operational_source_names(detail_meta):
    """Source DN names are optional; the cycle-level status is the approved shared view."""
    try:
        if not frappe.has_permission("Stock Entry", "read"):
            return False
        closed_field = detail_meta.get_field("custom_return_closed")
        permlevel = int(getattr(closed_field, "permlevel", 0) or 0)
        if permlevel:
            roles = set(frappe.get_roles(frappe.session.user)) | {"All"}
            stock_meta = frappe.get_meta("Stock Entry")
            field_read = any(
                str(getattr(permission, "role", "")) in roles
                and int(getattr(permission, "permlevel", 0) or 0) == permlevel
                and getattr(permission, "read", 0) in (1, True, "1", "Yes", "yes")
                for permission in stock_meta.permissions
            )
            if not field_read:
                return False
        return True
    except Exception:
        return False


def _warehouse_operational_lines(cycles):
    """Build a narrow cycle-site status view without granting Stock Entry access.

    Callers pass only cycles already loaded through permission-aware SIG Site Cycle
    reads and filtered to the user's routed scope (or the authorized all-project
    dashboard). This query exposes only the explicit closed flag for matching site
    DNs. Source DN names are retained only when the user already has Stock Entry and
    field-level read access.
    """
    try:
        detail_meta = frappe.get_meta("Stock Entry Detail")
        if not all(detail_meta.has_field(field) for field in ("custom_site", "custom_return_closed")):
            return [], False
        site_refs = sorted({
            str(cycle.get(field) or "").strip().upper()
            for cycle in cycles for field in ("site_key", "site")
            if str(cycle.get(field) or "").strip()
        })
        if not site_refs:
            return [], True
        lines = []
        for start in range(0, len(site_refs), 400):
            batch = site_refs[start:start + 400]
            placeholders = ", ".join(["%s"] * len(batch))
            rows = frappe.db.sql(
                """SELECT MIN(sed.custom_site) AS custom_site,
                          MAX(CASE WHEN IFNULL(sed.custom_return_closed, 0) = 1 THEN 1 ELSE 0 END)
                          AS custom_return_closed FROM `tabStock Entry Detail` sed
                   INNER JOIN `tabStock Entry` se ON se.name = sed.parent
                   WHERE sed.parenttype = %s
                     AND UPPER(IFNULL(sed.custom_site, '')) IN (""" + placeholders + """)
                     AND se.docstatus = %s AND se.is_return = %s
                     AND se.purpose IN (%s, %s)
                   GROUP BY UPPER(IFNULL(sed.custom_site, ''))""",
                tuple(["Stock Entry"] + batch + [1, 0, "Material Issue", "Material Transfer"]),
                as_dict=True,
            )
            lines.extend(rows)

        # Status is intentionally visible to cycle viewers, but DN identifiers
        # remain subject to Stock Entry record permissions and field access.
        if _can_read_operational_source_names(detail_meta):
            try:
                sources, source_read_complete = _all_visible(
                    "Stock Entry", ["name"],
                    filters={"docstatus": 1, "is_return": 0,
                             "purpose": ["in", ["Material Issue", "Material Transfer"]]},
                    order_by="posting_date asc, name asc", max_rows=MAX_ROWS)
                source_names = [str(row.get("name") or "") for row in sources if row.get("name")]
                if source_read_complete and source_names:
                    for start in range(0, len(source_names), 400):
                        batch = source_names[start:start + 400]
                        placeholders = ", ".join(["%s"] * len(batch))
                        source_lines = frappe.db.sql(
                            "SELECT sed.parent, sed.custom_site, sed.custom_return_closed "
                            "FROM `tabStock Entry Detail` sed "
                            "WHERE sed.parenttype = %s AND sed.parent IN (" + placeholders + ") "
                            "AND UPPER(IFNULL(sed.custom_site, '')) IN (" + ", ".join(["%s"] * len(site_refs)) + ")",
                            tuple(["Stock Entry"] + batch + site_refs),
                            as_dict=True,
                        )
                        lines.extend(source_lines)
            except Exception:
                # A failure in the optional identifier read must not hide the
                # cycle-level Open/Closed status already derived above.
                pass
        return lines, True
    except Exception:
        return [], False


def _warehouse_cycles_in_scope(rows, projects_by_name, memberships, can_see_all):
    """Limit the derived Warehouse status read to open cycles this caller may see."""
    open_rows = [row for row in rows if not _flag(row, "terminal")]
    if can_see_all:
        return open_rows
    visible_scopes = {str(m.get("scope") or "") for m in memberships}
    return [
        row for row in open_rows
        if _family((projects_by_name.get(str(row.get("erp_project") or "")) or {}).get("custom_family")
                   or row.get("project")) in visible_scopes
    ]


def _stage_tasks_for_user(cycles, memberships, all_project_scopes=False):
    """Project one queue row per logical task; never create one per employee."""
    tasks = []
    for cycle in cycles:
        if cycle.get("terminal"):
            continue
        if not all_project_scopes and not any(m.get("scope") == cycle.get("family") for m in memberships):
            continue
        for work in cycle.get("work") or []:
            role = work.get("owner_role") or "either"
            if not all_project_scopes and not _task_audience(cycle.get("family"), role, memberships):
                continue
            tasks.append({
                "id": cycle["name"] + "::" + str(work.get("key") or "work"),
                "kind": "stage", "cycle": cycle["name"], "site": cycle["site"],
                "project_id": cycle["project_id"], "ba": cycle["ba"],
                "alias": cycle["alias"], "wo": cycle["wo"],
                "family": cycle["family"], "cycle_stage": cycle["stage"],
                "stage": str(work.get("stage") or TASK_STAGE.get(str(work.get("key") or ""), cycle["stage"])),
                "key": str(work.get("key") or ""), "label": str(work.get("label") or "Work"),
                "owner_role": role, "owner_names": _owner_names(cycle.get("family"), role),
                "shared": True,
                "status": "Review due" if cycle.get("needs_review") else "Pending evidence",
                "needs_review": bool(cycle.get("needs_review")),
                "review_reason": str(cycle.get("delay_reason") or cycle.get("delay_note") or ""),
                "expected_evidence": str(work.get("closes_on") or ""),
            })

    return tasks


def _cycles_for_user(cycles, memberships, can_see_all=False):
    """Return open cycles for the user's confirmed project scope.

    PMs and coordinators share the cycle pool for their project family. Their
    cycle dashboard shows every open item in that family so the handoff is
    visible to the team. Warehouse users see only cycles with a warehouse
    action routed to them. System Managers can inspect the full visible pool.
    """
    visible = []
    if can_see_all:
        memberships = []
    scope_roles = {(str(item.get("scope") or ""), str(item.get("role") or ""))
                   for item in memberships}
    team_scopes = {scope for scope, role in scope_roles if role in ("pm", "co")}
    all_scopes = team_scopes | {scope for scope, role in scope_roles if role == "warehouse"}
    warehouse_only = bool(all_scopes) and not team_scopes

    for cycle in cycles:
        if cycle.get("terminal"):
            continue
        family = str(cycle.get("family") or "")
        if not can_see_all and family not in all_scopes:
            continue
        row = dict(cycle)
        work = list(cycle.get("work") or [])
        if warehouse_only:
            work = [item for item in work
                    if _task_audience(family, item.get("owner_role") or "either", memberships)]
            if not work:
                continue
        row["work"] = work
        visible.append(row)
    return visible


class _PlainText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.suppressed = [], 0
    def handle_starttag(self, tag, attrs):
        if tag.lower() in ("script", "style", "iframe", "object"):
            self.suppressed += 1
        elif tag.lower() in ("br", "p", "div", "li") and not self.suppressed:
            self.parts.append(" ")
    def handle_endtag(self, tag):
        if tag.lower() in ("script", "style", "iframe", "object") and self.suppressed:
            self.suppressed -= 1
        elif tag.lower() in ("p", "div", "li") and not self.suppressed:
            self.parts.append(" ")
    def handle_data(self, data):
        if not self.suppressed:
            self.parts.append(data)

def _plain_text(value):
    parser = _PlainText(); parser.feed(str(value or ""))
    return " ".join(unescape("".join(parser.parts)).split())

def _all_visible(doctype, fields, filters=None, order_by=None, max_rows=MAX_ROWS):
    rows, start = [], 0
    while start < max_rows:
        limit = min(PAGE_SIZE, max_rows - start)
        batch = frappe.get_list(doctype, filters=filters or {}, fields=fields,
                                order_by=order_by or "name asc", start=start, page_length=limit)
        rows.extend(batch)
        if len(batch) < limit:
            return rows, True
        start += len(batch)
    return rows, False

def _flag(row, field):
    return row.get(field) in (1, True, "1", "Yes", "yes")


def _can_close_cycles(user, memberships=None):
    memberships = memberships if memberships is not None else _current_ops_membership(user).get("memberships", [])
    roles = set(frappe.get_roles(user))
    admin = user != "Guest" and "System Manager" in roles
    ops_closer = any(m.get("role") in ("pm", "co") for m in memberships)
    field_ops = user != "Guest" and "SIG Field Ops" in roles and ops_closer
    return bool((admin or field_ops)
                and frappe.has_permission("SIG Site Cycle", "write")
                and frappe.has_permission("SIG Evidence Event", "create"))

def _family(value):
    normalized = str(value or "").strip().lower()
    if normalized in ("5g", "colocation", "colo"):
        return "Colocation"
    if normalized == "capex":
        return "Capex"
    if normalized == "incident":
        return "Incident"
    return ""


def _cycle_view(row, project, today, operational=None):
    """Derive the highest evidenced stage without inventing missing predecessors."""
    operational_visible = operational is not None
    operational = operational or {"applicable": False, "closed": False, "source_dns": []}
    closed = _flag(row, "terminal")
    has_prep = bool(row.get("mapping_at") or row.get("boq_at") or row.get("pimp_at")
                    or row.get("tcn_at") or row.get("mir_at") or row.get("first_visit_at"))
    prepared = bool(row.get("tcn_at") and row.get("mir_at") and row.get("first_visit_at"))
    if closed:
        stage = "Closed"
    elif row.get("invoice_at") or row.get("ho_at"):
        # Applicable handover evidence starts the final active Invoicing stage.
        # The PM-owned As-Built BOQ file remains a task within this stage.
        stage = "4 - Invoicing"
    elif prepared:
        stage = "3 - Handover"
    elif has_prep:
        stage = "2 - Plan"
    else:
        stage = "1 - Initial"
    work = []
    family = _family((project or {}).get("custom_family") or row.get("project"))
    if (not closed and not row.get("ho_at") and not row.get("asbuilt_at")
            and not row.get("invoice_at")
            and (row.get("mapping_at") or row.get("boq_at") or row.get("pimp_at"))):
        for key, label, field, role in (
            ("tcn", "TCN / MDT", "tcn_at", "co"),
            ("mir", "MIR", "mir_at", "co"),
            ("visit", "Visit", "first_visit_at", "pm"),
        ):
            if not row.get(field):
                work.append({"key": key, "label": label, "owner_role": role,
                             "closes_on": label + " evidence"})
    if not closed and prepared:
        for key, label, field, role in (
            ("tags", "Tags", "tags_at", "co"),
            ("photos", "Photos", "photos_at", "co"),
        ):
            if not row.get(field):
                work.append({"key": key, "label": label, "owner_role": role,
                             "closes_on": label + " evidence"})
        if operational.get("applicable") and not operational.get("closed"):
            work.append({"key": "asbuilt_operational", "label": "As-Built Operational",
                         "owner_role": "warehouse", "stage": "3 - Handover",
                         "closes_on": "Any linked site DN line custom_return_closed = 1",
                         "source_dns": operational.get("source_dns") or []})
    if not closed and (prepared or row.get("ho_at") or row.get("invoice_at")) and not row.get("asbuilt_at"):
        work.append({"key": "asbuilt", "label": "As-Built BOQ file", "owner_role": "pm",
                     "stage": "4 - Invoicing" if row.get("ho_at") or row.get("invoice_at") else "3 - Handover",
                     "closes_on": "ASBUILT_BOQ_GENERATED"})
    for item in work:
        item["owner_names"] = _owner_names(family, item.get("owner_role") or "either")
        item["shared"] = True
    delay = _flag(row, "is_delayed") or _flag(row, "is_aged") or bool(row.get("stage_due_on") and str(row.stage_due_on)[:10] < today)
    return {
        "name": row.name, "site": str(row.get("site_key") or row.get("site") or row.name),
        "project": str((project or {}).get("project_name") or row.get("project") or ""),
        "project_id": str((project or {}).get("name") or row.get("erp_project") or ""),
        "ba": str((project or {}).get("custom_agreement_no") or row.get("po") or ""),
        "alias": str(row.get("project_alias") or ""), "wo": str(row.get("wo") or ""),
        "family": family, "stage": stage, "legacy_stage": str(row.get("stage") or ""),
        "terminal": closed, "closed_on": str(row.get("closed_on") or ""),
        "closed_by": str(row.get("closed_by") or ""),
        "delay_reason": str(row.get("delay_reason") or ""),
        "delay_note": str(row.get("delay_note") or ""),
        "stage_due_on": str(row.get("stage_due_on") or ""),
        "last_event_at": str(row.get("last_event_at") or ""),
        "is_delayed": _flag(row, "is_delayed"), "is_aged": _flag(row, "is_aged"),
        "past_due": bool(row.get("stage_due_on") and str(row.stage_due_on)[:10] < today),
        "needs_review": bool(delay and not closed and work), "work": work,
        "events": {key: str(row.get(key) or "") for key in (
            "mapping_at", "boq_at", "tcn_at", "mir_at", "first_visit_at",
            "tags_at", "photos_at", "asbuilt_at", "warranty_at", "ho_at", "invoice_at")},
        "asbuilt_operational": {"visible": operational_visible,
                                 "applicable": bool(operational.get("applicable")) if operational_visible else None,
                                 "closed": bool(operational.get("closed")) if operational_visible else None,
                                 "source_dns": operational.get("source_dns") or [],
                                 "owner_names": _owner_names(family, "warehouse")},
    }

def _get_overview(all_project_cycles=False):
    """Read current ToDos, evidence-derived pipeline work and authorized cycles."""
    result = {"todos": [], "cycles": [], "projects": [], "stage_tasks": [],
                    "summary": {"todos": 0, "overdue_todos": 0, "cycles": 0, "delayed": 0,
                          "aged": 0, "past_due": 0, "closed": 0, "exceptions": 0,
                          "stages": {}, "stage_tasks": 0, "review_tasks": 0},
              "complete": {"todos": True, "cycles": True}, "todo_access": False,
              "can_close_cycles": False, "can_export": False, "scope": "assigned",
              "todo_create": False, "person": frappe.session.user,
              "task_route": {"ready": False, "issue": "", "memberships": []}}
    user, today = frappe.session.user, frappe.utils.nowdate()
    todos = []
    result["todo_create"] = bool(frappe.has_permission("ToDo", "create"))
    if frappe.has_permission("ToDo", "read"):
        result["todo_access"] = True
        todos, result["complete"]["todos"] = _all_visible(
            "ToDo", ["name", "description", "date", "status", "reference_type", "reference_name", "allocated_to"],
            filters={"allocated_to": user, "status": ["not in", ["Closed", "Cancelled"]]},
            order_by="date asc, modified desc", max_rows=5000)
        result["summary"]["todos"] = len(todos)
        result["summary"]["overdue_todos"] = sum(1 for row in todos if row.get("date") and str(row.date) < today)
        result["todos"] = [{**row, "description": _plain_text(row.get("description"))} for row in todos[:50]]
    if not frappe.has_permission("SIG Site Cycle", "read"):
        return result
    fields = ["name", "site", "site_key", "project", "erp_project", "project_stream", "po", "wo", "stage", "stage_due_on", "terminal",
              "closed_on", "closed_by", "is_delayed", "is_aged", "delay_reason", "delay_note", "last_event_at",
              "mapping_at", "boq_at", "pimp_at", "tcn_at", "mir_at", "first_visit_at", "tags_at", "photos_at",
              "asbuilt_at", "warranty_at", "ho_at", "invoice_at"]
    rows, result["complete"]["cycles"] = _all_visible(
        "SIG Site Cycle", fields, order_by="terminal asc, stage_due_on asc, last_event_at desc")
    projects_by_name = {}
    project_names = sorted({str(r.get("erp_project")) for r in rows if r.get("erp_project")})
    if project_names and frappe.has_permission("Project", "read"):
        projects, _ = _all_visible("Project", ["name", "project_name", "status", "custom_family", "custom_agreement_no"],
                                   filters={"name": ["in", project_names]}, order_by="project_name asc")
        projects_by_name = {p.name: p for p in projects}
    aliases_by_stream = {}
    stream_names = sorted({str(r.get("project_stream")) for r in rows if r.get("project_stream")})
    if stream_names and frappe.db.exists("DocType", "SIG Project Stream") and frappe.has_permission("SIG Project Stream", "read"):
        streams, _ = _all_visible("SIG Project Stream", ["name", "alias"],
                                  filters={"name": ["in", stream_names]}, max_rows=5000)
        aliases_by_stream = {str(stream.name): str(stream.get("alias") or "") for stream in streams}
    result["task_route"] = _current_ops_membership(user)
    route_memberships = result["task_route"].get("memberships") or []
    can_see_all_cycles = all_project_cycles or (
        user != "Guest" and "System Manager" in set(frappe.get_roles(user))
    )
    warehouse_cycles = _warehouse_cycles_in_scope(
        rows, projects_by_name, route_memberships, can_see_all_cycles)
    counts = result["summary"]
    warehouse_lines, warehouse_complete = _warehouse_operational_lines(warehouse_cycles)
    result["complete"]["asbuilt_operational"] = warehouse_complete
    result["operational_access"] = (
        "available" if warehouse_complete is True
        else "restricted" if warehouse_complete is None
        else "unavailable"
    )
    for row in rows:
        project_key = str(row.get("erp_project") or "")
        project = projects_by_name.get(project_key)
        row["project_alias"] = aliases_by_stream.get(str(row.get("project_stream") or ""), "")
        operational = _asbuilt_operational_status(row, warehouse_lines) if warehouse_complete else None
        cycle = _cycle_view(row, project, today, operational)
        result["cycles"].append(cycle)
        if cycle["terminal"]:
            counts["closed"] += 1
        else:
            counts["cycles"] += 1
            counts["delayed"] += int(cycle["is_delayed"])
            counts["aged"] += int(cycle["is_aged"])
            counts["past_due"] += int(cycle["past_due"])
            counts["exceptions"] += int(cycle["needs_review"])
        counts["stages"][cycle["stage"]] = counts["stages"].get(cycle["stage"], 0) + 1
    result["scope"] = "all" if can_see_all_cycles else "assigned"
    result["can_export"] = bool(frappe.has_permission("SIG Site Cycle", "export"))
    result["cycles"] = _cycles_for_user(result["cycles"], route_memberships, can_see_all_cycles)
    result["cycles"].sort(key=lambda r: (not r["needs_review"], r["stage_due_on"] or "9999-12-31", r["site"].lower()))

    # The client receives only the open, in-scope cycle pool. Recompute these
    # totals after routing so the Workbench and the cycles subpage agree.
    counts.update({"cycles": 0, "delayed": 0, "aged": 0, "past_due": 0,
                   "exceptions": 0, "closed": 0, "stages": {}})
    for cycle in result["cycles"]:
        counts["cycles"] += 1
        counts["delayed"] += int(cycle["is_delayed"])
        counts["aged"] += int(cycle["is_aged"])
        counts["past_due"] += int(cycle["past_due"])
        counts["exceptions"] += int(cycle["needs_review"])
        counts["stages"][cycle["stage"]] = counts["stages"].get(cycle["stage"], 0) + 1

    result["can_close_cycles"] = _can_close_cycles(user, route_memberships)
    if all_project_cycles:
        result["stage_tasks"] = _stage_tasks_for_user(
            result["cycles"], result["task_route"]["memberships"], all_project_scopes=True)
        result["summary"]["stage_tasks"] = sum(1 for task in result["stage_tasks"] if task["kind"] == "stage")
        result["summary"]["review_tasks"] = sum(1 for task in result["stage_tasks"]
                                                if task["kind"] == "stage" and task["needs_review"])
    elif result["task_route"]["ready"]:
        result["stage_tasks"] = _stage_tasks_for_user(result["cycles"], result["task_route"]["memberships"])
        result["summary"]["stage_tasks"] = sum(1 for task in result["stage_tasks"] if task["kind"] == "stage")
        result["summary"]["review_tasks"] = sum(1 for task in result["stage_tasks"]
                                                if task["kind"] == "stage" and task["needs_review"])
    result["todos_has_more"] = len(todos) > len(result["todos"]) if result["todo_access"] else False
    return result


@frappe.whitelist()
def get_overview():
    """Return the signed-in user's ToDos and assigned Operations scope."""
    return _get_overview()


def _require_cycle_dashboard_access():
    user = frappe.session.user
    roles = set(frappe.get_roles(user)) if user != "Guest" else set()
    if (user == "Guest" or not roles.intersection(SITE_CYCLE_DASHBOARD_ROLES)
            or not frappe.has_permission("SIG Site Cycle", "read")):
        frappe.throw("You need an Operations page role and SIG Site Cycle read access to view all project cycles.",
                     frappe.PermissionError)


@frappe.whitelist()
def get_cycle_dashboard():
    """Return every open cycle visible to the current user across project scopes."""
    _require_cycle_dashboard_access()
    data = _get_overview(all_project_cycles=True)
    cycles = [cycle for cycle in data.get("cycles") or [] if not cycle.get("terminal")]
    stages = {stage: sum(1 for cycle in cycles if cycle.get("stage") == stage)
              for stage in ("1 - Initial", "2 - Plan", "3 - Handover", "4 - Invoicing")}
    return {
        "cycles": cycles,
        "route": data.get("task_route") or {"ready": False, "issue": "", "memberships": []},
        "scope": "all",
        "can_export": bool(frappe.has_permission("SIG Site Cycle", "export")),
        "complete": data.get("complete", {}).get("cycles", True),
        "summary": {
            "cycles": len(cycles),
            "open_items": sum(len(cycle.get("work") or []) for cycle in cycles),
            "review_due": sum(1 for cycle in cycles if cycle.get("needs_review")),
            "stages": stages,
        },
    }

@frappe.whitelist()
def export_workbook(filter="open", scope="assigned"):
    """Export the opener's queue, or the all-project cycle dashboard when requested."""
    if not frappe.has_permission("SIG Site Cycle", "read") or not frappe.has_permission("SIG Site Cycle", "export"):
        frappe.throw("You need SIG Site Cycle read and export permission to download this workbook.",
                     frappe.PermissionError)
    if scope == "all":
        _require_cycle_dashboard_access()
        if filter != "cycles":
            frappe.throw("The all-project export is available for the Site Cycles report only.")
        data = _get_overview(all_project_cycles=True)
    else:
        data = get_overview()
    if not data.get("complete", {}).get("cycles"):
        frappe.throw("The visible cycle read reached its safety limit. Refine the source data before exporting.")
    operational_complete = data.get("complete", {}).get("asbuilt_operational") is True
    operational_access = data.get("operational_access") or "unavailable"
    all_cycles = list(data.get("cycles") or [])
    if filter == "cycles" or filter == "open":
        cycles = [c for c in all_cycles if not c.get("terminal")]
        tasks = [t for t in data.get("stage_tasks") or [] if t.get("cycle") in {c.get("name") for c in cycles}]
    else:
        tasks = list(data.get("stage_tasks") or [])
        if filter == "review":
            tasks = [t for t in tasks if t.get("needs_review") or t.get("kind") == "review"]
        elif filter.startswith("stage:"):
            stage = filter[6:]
            tasks = [t for t in tasks if (t.get("kind") == "closure" if stage == "Cycle closure" else t.get("stage") == stage and t.get("kind") != "closure")]
        else:
            tasks = [t for t in tasks if t.get("kind") != "review"]
        cycle_names = {t.get("cycle") for t in tasks}
        cycles = [c for c in all_cycles if c.get("name") in cycle_names and not c.get("terminal")]

    wb = Workbook()
    ws = wb.active
    ws.title = "Cycles"
    headers = ["Cycle ID", "Site", "WO", "BA", "Project alias", "ERP Project", "Scope",
               "Stage", "Status", "Pending work", "Pending item owners", "Review due", "Delayed", "Aged",
               "Due date", "Delay reason", "Delay note", "Closed on", "Closed by",
               "Mapping at", "BOQ at", "TCN at", "MIR at", "Visit at",
               "Tags at", "Photos at", "As-Built BOQ at", "As-Built Operational status",
               "As-Built DN source(s) (if accessible)", "Warranty at", "Handover at", "Invoice at"]
    ws.append(headers)
    stage_ws = wb.create_sheet("Stage tasks")
    stage_ws.append(["Task ID", "Cycle ID", "Site", "WO", "BA", "Project alias",
                     "Task stage", "Cycle stage", "Task", "Owner(s)", "Status",
                     "Expected evidence", "Review reason"])
    cycle_by_name = {c.get("name"): c for c in all_cycles}
    for cycle in cycles:
        events = cycle.get("events") or {}
        ws.append([cycle.get("name"), cycle.get("site"), cycle.get("wo"), cycle.get("ba"),
                   cycle.get("alias"), cycle.get("project_id"), cycle.get("family"),
                   cycle.get("stage"), "Closed" if cycle.get("terminal") else "Open",
                   "; ".join(w.get("label", "") for w in cycle.get("work") or []),
                   "; ".join(
                       f"{w.get('label', 'Work')}: {' / '.join(w.get('owner_names') or [])}"
                       for w in cycle.get("work") or []),
                   bool(cycle.get("needs_review")), bool(cycle.get("is_delayed")), bool(cycle.get("is_aged")),
                   cycle.get("stage_due_on"), cycle.get("delay_reason"), cycle.get("delay_note"),
                   cycle.get("closed_on"), cycle.get("closed_by")] +
                  [events.get(field) for field in ("mapping_at", "boq_at", "tcn_at", "mir_at",
                   "first_visit_at", "tags_at", "photos_at", "asbuilt_at")])
        operational = cycle.get("asbuilt_operational") or {}
        if not operational_complete:
            operational_status = (
                "Unavailable to this account" if operational_access == "restricted"
                else "Coverage could not be verified"
            )
            source_dns = ""
        elif not operational.get("visible"):
            operational_status = "Coverage could not be verified"
            source_dns = ""
        elif not operational.get("applicable"):
            operational_status = "No matching DN"
            source_dns = ""
        else:
            operational_status = "Closed" if operational.get("closed") else "Open"
            source_dns = "; ".join(operational.get("source_dns") or [])
        ws.cell(ws.max_row, 28, operational_status)
        ws.cell(ws.max_row, 29, source_dns)
        ws.cell(ws.max_row, 30, events.get("warranty_at"))
        ws.cell(ws.max_row, 31, events.get("ho_at"))
        ws.cell(ws.max_row, 32, events.get("invoice_at"))
    for task in tasks:
        cycle = cycle_by_name.get(task.get("cycle"))
        if not cycle:
            continue
        audience = " / ".join(task.get("owner_names") or [])
        stage_ws.append([task.get("id"), cycle.get("name"), cycle.get("site"), cycle.get("wo"),
                         cycle.get("ba"), cycle.get("alias"), task.get("stage"), cycle.get("stage"),
                         task.get("label"), audience, task.get("status"),
                         task.get("expected_evidence"), task.get("review_reason")])
    if not operational_complete:
        note_ws = wb.create_sheet("Export notes")
        note_ws.append(["Data coverage", "Details"])
        note_ws.append([
            "As-Built Operational",
            "Stock Entry coverage was not available to this account. Cycle and task rows include accessible data; "
            "As-Built Operational statuses and related Warehouse tasks may be missing."
        ])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173F62")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    output = BytesIO()
    wb.save(output)
    frappe.local.response.filename = "operations-cycle-stage.xlsx"
    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.type = "download"

@frappe.whitelist()
def close_cycles(cycles, reason):
    """Write explicit SITE_CLOSED evidence for selected cycles after role and permission checks."""
    user = frappe.session.user
    if not _can_close_cycles(user):
        frappe.throw("Only Field Ops users may close site cycles.", frappe.PermissionError)
    try:
        names = cycles if isinstance(cycles, list) else json.loads(cycles or "[]")
    except Exception:
        frappe.throw("Invalid cycle selection.")
    names = list(dict.fromkeys(str(n).strip() for n in names if str(n).strip()))
    reason = str(reason or "").strip().replace("|", "/")
    if not names or len(names) > 50 or not reason or len(reason) > 500:
        frappe.throw("Select 1–50 cycles and enter a closure reason up to 500 characters.")
    docs = []
    for name in names:
        if not frappe.db.exists("SIG Site Cycle", name):
            frappe.throw("A selected site cycle no longer exists. Refresh and review the selection.")
        if not frappe.has_permission("SIG Site Cycle", "write", name):
            frappe.throw("You cannot close every selected cycle. No changes were made.", frappe.PermissionError)
        doc = frappe.get_doc("SIG Site Cycle", name)
        if _flag(doc, "terminal"):
            continue
        docs.append(doc)
    now = frappe.utils.now()
    stamp = now.replace("-", "").replace(" ", "").replace(":", "")
    safe_user = user.replace("@", "_").replace(".", "_")
    batch = "MANUAL-" + safe_user + "-" + stamp
    closed = []
    for seq, cycle in enumerate(docs, start=1):
        cycle.terminal = 1
        cycle.closed_on = now[:10]
        cycle.closed_by = user
        cycle.stage = "5 - Closed"
        cycle.stage_anchor_at = None
        cycle.stage_due_on = None
        cycle.is_delayed = 0
        cycle.last_event_at = now
        cycle.last_event_code = "SITE_CLOSED"
        cycle.save()
        event_id = batch + "-" + str(seq).zfill(6)
        event = frappe.get_doc({
            "doctype": "SIG Evidence Event", "event_id": event_id,
            "schema_version": "EV1", "event_code": "SITE_CLOSED",
            "occurred_at": now, "recorded_at": now, "site_id": cycle.site_id,
            "site_key": cycle.site_key, "project": cycle.project, "po": cycle.po,
            "wo": cycle.wo, "artifact_ref": "MANUAL_CLOSE|" + user + "|" + now + "|" + reason,
            "batch_id": batch, "seq": seq, "source_sub": "operations_workbench",
            "source_workbook": "ERP", "operator": user, "machine": "neosys",
            "core_version": "erp", "cycle": cycle.name,
        })
        event.insert()
        closed.append(cycle.name)
    return {"result": "ok", "closed": closed, "already_closed": len(names) - len(docs),
            "batch": batch, "operator": user, "reason": reason}

