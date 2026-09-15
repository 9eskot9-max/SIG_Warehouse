"""Turn the Stage 5 As-Built acceptance JSON into an owner triage pack.

This is a read-only review artifact. It does not alter the acceptance JSON or
write to ERP; it only groups non-MATCH rows by the next evidence/action.
"""

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook


DEFAULT_INPUT = Path(r"C:\Users\legen\Downloads\wh\outputs\WH_Stage5_AsBuilt_Acceptance_2026-09-13.json")
DEFAULT_OUTPUT = Path(r"C:\Users\legen\Downloads\wh\outputs\WH_Stage5_AsBuilt_Triage_2026-09-13.xlsx")


ACTION_BY_STATUS = {
    "MISMATCH": "Warehouse + ERP: classify posting/source discrepancy",
    "SITE_ALLOCATION_REVIEW": "ERP: assign source line to a site or document as warehouse-only",
    "EXCEPTION": "Warehouse + ERP: resolve allocation/site exception before acceptance",
    "WH_ONLY": "Warehouse: find missing ERP movement or confirm historical-only row",
    "ERP_ONLY": "ERP: find unmatched WH source or confirm ERP-only movement",
}


def triage_rows(rows):
    out = []
    for row in rows:
        status = row.get("status")
        if status == "MATCH":
            continue
        enriched = dict(row)
        enriched["recommended_action"] = ACTION_BY_STATUS.get(status, "Review")
        enriched["absolute_difference"] = abs(float(row.get("difference_erp_minus_wh") or 0))
        out.append(enriched)
    return sorted(out, key=lambda row: (-row["absolute_difference"], str(row.get("site") or ""), row["item_code"]))


def build_pack(input_path, output_path):
    report = json.loads(Path(input_path).read_text(encoding="utf-8"))
    rows = triage_rows(report.get("rows", []))
    counts = Counter(row["status"] for row in rows)
    by_site = defaultdict(lambda: {"rows": 0, "abs_difference": 0.0})
    for row in rows:
        site = row.get("site") or "<NO SITE>"
        by_site[site]["rows"] += 1
        by_site[site]["abs_difference"] += row["absolute_difference"]

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["SIG Stage 5 As-Built triage", None])
    summary.append(["Source acceptance", str(input_path)])
    summary.append(["Acceptance status", report.get("acceptance", {}).get("status")])
    summary.append(["As of", report.get("as_of")])
    summary.append([])
    summary.append(["Status", "Rows", "Recommended action"])
    for status in ("MISMATCH", "SITE_ALLOCATION_REVIEW", "EXCEPTION", "WH_ONLY", "ERP_ONLY"):
        summary.append([status, counts.get(status, 0), ACTION_BY_STATUS[status]])
    summary.append([])
    summary.append(["Site", "Non-match rows", "Absolute quantity difference"])
    for site, values in sorted(by_site.items(), key=lambda pair: -pair[1]["abs_difference"]):
        summary.append([site, values["rows"], round(values["abs_difference"], 6)])

    detail = wb.create_sheet("Non-Match Triage")
    headers = [
        "status", "recommended_action", "site", "item_code", "wh_issued", "wh_returned",
        "wh_custody", "wh_expected_as_built", "erp_issued", "erp_as_built",
        "difference_erp_minus_wh", "absolute_difference", "erp_unallocated_item_qty", "exceptions",
    ]
    detail.append(headers)
    for row in rows:
        detail.append([row.get(header) for header in headers])
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions

    for status in ("MISMATCH", "SITE_ALLOCATION_REVIEW", "EXCEPTION", "WH_ONLY", "ERP_ONLY"):
        sheet = wb.create_sheet(status[:31])
        sheet.append(headers)
        for row in rows:
            if row["status"] == status:
                sheet.append([row.get(header) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {
        "result": "generated",
        "acceptance": report.get("acceptance", {}).get("status"),
        "non_match_rows": len(rows),
        "status_counts": dict(counts),
        "xlsx": str(output_path),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(json.dumps(build_pack(args.input, args.output), indent=2))


if __name__ == "__main__":
    main()
