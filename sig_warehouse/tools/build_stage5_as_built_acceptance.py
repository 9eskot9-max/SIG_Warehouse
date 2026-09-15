"""Build a read-only Stage 5 As-Built acceptance pack.

The pack compares the frozen WH closure workbook with submitted ERP Stock
Entry lines.  It never creates or updates an ERP document.  Differences are
classified for review; matching totals are not treated as acceptance by
themselves.
"""

import argparse
import csv
import hashlib
import json
import os
import sys
import urllib.parse
import urllib.request
import urllib.error
import time
from collections import defaultdict
from datetime import date
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover - site runtime has openpyxl in the bundle
    Workbook = load_workbook = None


APP_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = Path(__file__).resolve().parents[3]
PACKAGE_ROOT = Path(__file__).resolve().parents[2]
if str(PACKAGE_ROOT) not in sys.path:
    sys.path.insert(0, str(PACKAGE_ROOT))

try:
    from sig_warehouse.sig_warehouse.stage5_as_built import build_as_built_rows
except ImportError:  # pragma: no cover - direct script invocation fallback
    build_as_built_rows = None


EPS = 1e-6
AS_OF = "2026-09-13"
DEFAULT_WH = Path(r"C:\Users\legen\Downloads\wh\outputs\WH_Closure_Package_2026.xlsx")
DEFAULT_JSON = Path(r"C:\Users\legen\Downloads\wh\outputs\WH_Stage5_AsBuilt_Acceptance_2026-09-13.json")
DEFAULT_XLSX = Path(r"C:\Users\legen\Downloads\wh\outputs\WH_Stage5_AsBuilt_Acceptance_2026-09-13.xlsx")


def load_env(path):
    values = {}
    for raw in Path(path).read_text(encoding="utf-8").splitlines():
        if "=" not in raw or raw.lstrip().startswith("#"):
            continue
        key, value = raw.split("=", 1)
        values[key.strip()] = value.strip().strip('"')
    return values


class ErpApi:
    def __init__(self, env):
        self.site = env["SIG_ERPNEXT_SITE"].rstrip("/")
        token = f"{env['SIG_ERPNEXT_API_KEY']}:{env['SIG_ERPNEXT_API_SECRET']}"
        self.headers = {"Authorization": f"token {token}", "Accept": "application/json"}

    def get(self, doctype, params):
        query = urllib.parse.urlencode({k: json.dumps(v) if isinstance(v, (list, dict)) else str(v)
                                        for k, v in params.items()})
        if doctype.startswith("DocType/"):
            resource = "DocType/" + urllib.parse.quote(doctype.split("/", 1)[1], safe="")
        elif "/" in doctype:
            parent, name = doctype.split("/", 1)
            resource = urllib.parse.quote(parent, safe="") + "/" + urllib.parse.quote(name, safe="")
        else:
            resource = urllib.parse.quote(doctype, safe="")
        url = f"{self.site}/api/resource/{resource}?{query}"
        request = urllib.request.Request(url, headers=self.headers)
        with urllib.request.urlopen(request, timeout=60) as response:
            if response.status != 200:
                raise RuntimeError(f"{doctype}: HTTP {response.status}")
            body = json.loads(response.read().decode("utf-8"))
        return body.get("data", [])

    def get_doc(self, doctype):
        if doctype.startswith("DocType/"):
            resource = "DocType/" + urllib.parse.quote(doctype.split("/", 1)[1], safe="")
        elif "/" in doctype:
            parent, name = doctype.split("/", 1)
            resource = urllib.parse.quote(parent, safe="") + "/" + urllib.parse.quote(name, safe="")
        else:
            resource = urllib.parse.quote(doctype, safe="")
        request = urllib.request.Request(f"{self.site}/api/resource/{resource}", headers=self.headers)
        for attempt in range(3):
            try:
                with urllib.request.urlopen(request, timeout=60) as response:
                    if response.status != 200:
                        raise RuntimeError(f"{doctype}: HTTP {response.status}")
                    return json.loads(response.read().decode("utf-8")).get("data", {})
            except urllib.error.HTTPError:
                if attempt == 2:
                    raise
                time.sleep(1.0 * (attempt + 1))

    def fetch_all(self, doctype, filters, fields, order_by=None):
        result = []
        start = 0
        while True:
            params = {
                "filters": filters,
                "fields": fields,
                "limit_start": start,
                "limit_page_length": 500,
            }
            if order_by:
                params["order_by"] = order_by
            page = self.get(doctype, params)
            result.extend(page)
            if len(page) < 500:
                return result
            start += 500


def number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def header_row(ws, required):
    wanted = {str(value).strip() for value in required}
    for row_number, values in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), 1):
        headers = {str(value).strip() for value in values if value is not None}
        if wanted.issubset(headers):
            return row_number, {str(value).strip(): index for index, value in enumerate(values)}
    raise ValueError(f"Could not find headers {sorted(wanted)} in {ws.title}")


def read_wh_snapshot(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    issued = defaultdict(float)
    returned = defaultdict(float)
    custody = defaultdict(float)

    ws = wb["Outcome Txns Full (Jan-13Sep)"]
    row_number, cols = header_row(ws, {"Item Code", "Qty", "Site ID"})
    for values in ws.iter_rows(min_row=row_number + 1, values_only=True):
        item, site, qty = str(values[cols["Item Code"]] or "").strip(), str(values[cols["Site ID"]] or "").strip(), number(values[cols["Qty"]])
        if item and qty:
            issued[(site or None, item)] += qty

    ws = wb["Returns Full (Mar-13Sep)"]
    row_number, cols = header_row(ws, {"Item Code", "Qty", "Site ID"})
    for values in ws.iter_rows(min_row=row_number + 1, values_only=True):
        item, site, qty = str(values[cols["Item Code"]] or "").strip(), str(values[cols["Site ID"]] or "").strip(), number(values[cols["Qty"]])
        if item and qty:
            returned[(site or None, item)] += qty

    ws = wb["In-Hand Activity"]
    row_number, cols = header_row(ws, {"Event Type", "Item Code", "Qty", "Source DN ID"})
    deltas = {"OPEN": 1.0, "REASSIGN_IN": 1.0, "RETURN": -1.0, "USE": -1.0, "REASSIGN_OUT": -1.0}
    dn_sites = {}
    out_csv = Path(__file__).resolve().parents[3] / "scripts" / "wh_stock" / "out_confirmed_with_comment.csv"
    if out_csv.exists():
        with out_csv.open(encoding="utf-8-sig", newline="") as stream:
            for row in csv.DictReader(stream):
                dn = (row.get("dn") or "").strip()
                site = (row.get("site") or "").strip()
                if dn and site:
                    dn_sites.setdefault(dn, site)
    for values in ws.iter_rows(min_row=row_number + 1, values_only=True):
        event_type = str(values[cols["Event Type"]] or "").strip().upper()
        item = str(values[cols["Item Code"]] or "").strip()
        source_dn = str(values[cols["Source DN ID"]] or "").strip()
        site = dn_sites.get(source_dn)
        qty = number(values[cols["Qty"]])
        if event_type in deltas and item and qty:
            custody[(site, item)] += deltas[event_type] * qty

    return issued, returned, custody


def read_erp(api):
    source_filters = [
        ["docstatus", "=", 1],
        ["is_return", "=", 0],
        ["purpose", "in", ["Material Issue", "Material Transfer"]],
        ["posting_date", ">=", "2026-01-01"],
        ["posting_date", "<=", AS_OF],
    ]
    sources = api.fetch_all("Stock Entry", source_filters, ["name", "posting_date", "purpose"], "posting_date asc, creation asc")
    source_names = [row["name"] for row in sources]
    if not source_names:
        return sources, []

    meta = api.get_doc("DocType/Stock Entry Detail")
    available = {row["fieldname"] for row in meta.get("fields", [])}
    desired = [
        "name", "parent", "parenttype", "item_code", "qty", "uom", "custom_site", "project",
        "custom_row_key", "custom_qty_returned", "custom_qty_custody", "custom_return_disposition",
    ]
    fields = [field for field in desired if field in available]
    lines = []
    try:
        for offset in range(0, len(source_names), 100):
            chunk = source_names[offset:offset + 100]
            filters = [["parent", "in", chunk], ["parenttype", "=", "Stock Entry"]]
            lines.extend(api.fetch_all("Stock Entry Detail", filters, fields, "parent asc, idx asc"))
    except urllib.error.HTTPError as error:
        if error.code != 403:
            raise
        # Some named API users can read Stock Entry but are intentionally not
        # granted list permission on its child table.  Full parent reads are
        # still allowed, so use the native child payload as a read-only
        # fallback rather than widening permissions just for this report.
        from concurrent.futures import ThreadPoolExecutor

        def fetch_parent(name):
            document = api.get_doc(f"Stock Entry/{name}")
            return document.get("items", [])

        with ThreadPoolExecutor(max_workers=8) as pool:
            for child_rows in pool.map(fetch_parent, source_names):
                for row in child_rows:
                    lines.append(row)
    return sources, lines


def build_report(wh_path, env_path):
    api = ErpApi(load_env(env_path))
    wh_issued, wh_returned, wh_custody = read_wh_snapshot(wh_path)
    sources, lines = read_erp(api)
    source_by_name = {row["name"]: type("Source", (), row) for row in sources}
    line_objects = [type("Line", (), row) for row in lines]
    if build_as_built_rows is None:
        raise RuntimeError("Stage 5 As-Built module could not be imported")
    erp_rows = build_as_built_rows([type("Source", (), row) for row in sources], line_objects)

    erp_issue = defaultdict(float)
    erp_as_built = defaultdict(float)
    erp_exceptions = defaultdict(list)
    for row in erp_rows:
        key = (row.get("site") or None, row["item_code"])
        erp_issue[key] += row["issued_qty"]
        erp_as_built[key] += row["as_built_qty"]
        if row["exceptions"]:
            erp_exceptions[key].extend(row["exceptions"])

    erp_unallocated_by_item = defaultdict(float)
    for (site, item), qty in erp_issue.items():
        if site is None:
            erp_unallocated_by_item[item] += qty

    wh_keys = set(wh_issued) | set(wh_returned) | set(wh_custody)
    erp_keys = set(erp_issue) | set(erp_as_built)
    rows = []
    for key in sorted(wh_keys | erp_keys, key=lambda value: (str(value[0] or ""), str(value[1]))):
        site, item = key
        wh_expected = wh_issued.get(key, 0.0) - wh_returned.get(key, 0.0) - wh_custody.get(key, 0.0)
        erp_value = erp_as_built.get(key, 0.0)
        difference = erp_value - wh_expected
        flags = sorted(set(erp_exceptions.get(key, [])))
        if wh_expected < -EPS:
            flags.append("WH_ALLOCATION_EXCEEDS_ISSUED")
        unallocated_qty = erp_unallocated_by_item.get(item, 0.0)
        if flags:
            status = "EXCEPTION"
        elif unallocated_qty and (site is None or key not in erp_keys):
            # Do not pretend an untagged ERP line is a clean WH-only or
            # ERP-only delta.  It may belong to one of several WH sites and
            # needs source-line allocation evidence before acceptance.
            status = "SITE_ALLOCATION_REVIEW"
        elif key not in erp_keys:
            status = "WH_ONLY"
        elif key not in wh_keys:
            status = "ERP_ONLY"
        elif abs(difference) <= 1e-4:
            status = "MATCH"
        else:
            status = "MISMATCH"
        rows.append({
            "site": site, "item_code": item, "wh_issued": round(wh_issued.get(key, 0.0), 6),
            "wh_returned": round(wh_returned.get(key, 0.0), 6),
            "wh_custody": round(wh_custody.get(key, 0.0), 6),
            "wh_expected_as_built": round(wh_expected, 6),
            "erp_issued": round(erp_issue.get(key, 0.0), 6),
            "erp_as_built": round(erp_value, 6), "difference_erp_minus_wh": round(difference, 6),
            "erp_unallocated_item_qty": round(unallocated_qty, 6),
            "status": status, "exceptions": ";".join(sorted(set(flags))),
        })

    counts = defaultdict(int)
    for row in rows:
        counts[row["status"]] += 1
    return {
        "result": "generated", "as_of": AS_OF,
        "wh_snapshot": {"path": str(wh_path), "sha256": sha256(wh_path)},
        "erp": {"submitted_stock_entries": len(sources), "stock_entry_detail_lines": len(lines)},
        "summary": {"row_count": len(rows), "status_counts": dict(sorted(counts.items()))},
        "acceptance": {
            "status": "NOT_ACCEPTED",
            "reason": "The first comparison contains exceptions and classified deltas; review and classify each non-MATCH row before any cutover gate.",
            "blocking_statuses": [status for status in ("EXCEPTION", "MISMATCH", "SITE_ALLOCATION_REVIEW", "WH_ONLY", "ERP_ONLY") if counts.get(status)],
        },
        "rows": rows,
    }


def write_xlsx(report, path):
    if Workbook is None:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["SIG Stage 5 As-Built acceptance pack", None])
    ws.append(["As of", report["as_of"]])
    ws.append(["WH snapshot SHA-256", report["wh_snapshot"]["sha256"]])
    ws.append(["ERP submitted Stock Entries", report["erp"]["submitted_stock_entries"]])
    ws.append(["ERP detail lines", report["erp"]["stock_entry_detail_lines"]])
    ws.append([])
    for status, count in report["summary"]["status_counts"].items():
        ws.append([status, count])
    detail = wb.create_sheet("As-Built Delta")
    headers = list(report["rows"][0]) if report["rows"] else ["site", "item_code", "status"]
    detail.append(headers)
    for row in report["rows"]:
        detail.append([row.get(header) for header in headers])
    detail.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return True


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--wh-snapshot", type=Path, default=DEFAULT_WH)
    parser.add_argument("--env", type=Path, default=REPO_ROOT / ".secrets" / "sig_erpnext_api.env")
    parser.add_argument("--json-out", type=Path, default=DEFAULT_JSON)
    parser.add_argument("--xlsx-out", type=Path, default=DEFAULT_XLSX)
    args = parser.parse_args()
    report = build_report(args.wh_snapshot, args.env)
    args.json_out.parent.mkdir(parents=True, exist_ok=True)
    args.json_out.write_text(json.dumps(report, indent=2), encoding="utf-8")
    write_xlsx(report, args.xlsx_out)
    print(json.dumps({"result": report["result"], "acceptance": report["acceptance"], "summary": report["summary"], "json": str(args.json_out), "xlsx": str(args.xlsx_out)}, indent=2))


if __name__ == "__main__":
    main()
